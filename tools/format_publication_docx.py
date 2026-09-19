#!/usr/bin/env python3
"""Apply the project's restrained, black-only publication style to DOCX files."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def set_run_style(run, size=11, bold=None):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def format_document(path: Path):
    doc = Document(path)
    is_paper = path.name.startswith("De-la-IA-")

    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = "Aptos"
            style.font.color.rgb = RGBColor(0, 0, 0)

    abstract_mode = False
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""

        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_style(run, 16, True)
        elif index == 1 and is_paper:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_style(run, 12, True)
        elif index in (2, 3) and is_paper:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_style(run, 10, index == 2)
        elif style_name.startswith("Heading"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            level = int(style_name.split()[-1]) if style_name.split()[-1].isdigit() else 2
            size = 13 if level == 1 else 11
            for run in paragraph.runs:
                set_run_style(run, size, True)
            abstract_mode = text.lower() == "resumen"
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if style_name == "Source Code" else WD_ALIGN_PARAGRAPH.JUSTIFY
            size = 10 if abstract_mode else 11
            for run in paragraph.runs:
                set_run_style(run, size)
                if style_name == "Source Code":
                    run.font.name = "Aptos Mono"
                    run.font.size = Pt(9)
            if text.startswith("Palabras clave:"):
                abstract_mode = False

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_style(run, 9)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(54)
        section.left_margin = section.right_margin = Pt(58)

    doc.save(path)


if __name__ == "__main__":
    for filename in sys.argv[1:]:
        format_document(Path(filename))
